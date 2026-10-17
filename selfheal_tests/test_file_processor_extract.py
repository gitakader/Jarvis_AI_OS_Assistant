"""Offline checks for file_processor.extract_text returning content inline.

Root cause being guarded: extract_text used to reply only "Saved: file.txt",
so the live model answered questions about a PDF/Word doc from imagination
(2026-10-15: JARVIS invented two different phone numbers for Abdul_Kader_CV.pdf
because the extracted text never reached it)."""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import actions.file_processor as FP
from core.action_loader import _validate

PASS = 0
FAIL = 0

def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS  {name}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}  {detail}")

def _make_pdf(text: str) -> bytes:
    """Build a minimal one-page PDF with a correct xref table (no deps)."""
    content = f"BT /F1 14 Tf 72 700 Td ({text}) Tj ET"
    bodies = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
        + content.encode() + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(bodies, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(bodies)+1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(bodies)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)

def main():
    tmp = Path(tempfile.mkdtemp(prefix="fp_extract_"))

    # ── PDF: content must come back inline, file still saved ────────────────
    pdf = tmp / "cv.pdf"
    pdf.write_bytes(_make_pdf("CV Phone 01717-288 187 Email mukul.kader@gmail.com"))
    out = FP.file_processor({"file_path": str(pdf), "action": "extract_text"})
    check("pdf: phone inline", "01717-288 187" in out, out[:200])
    check("pdf: email inline", "mukul.kader@gmail.com" in out, out[:200])
    check("pdf: full-text note", "Full text (" in out and "saved to" in out, out[-160:])
    saved = tmp / "cv_text.txt"
    check("pdf: txt file written", saved.exists() and
          "01717-288 187" in saved.read_text(encoding="utf-8"))

    # ── PDF: long text is truncated, not dropped ────────────────────────────
    big = tmp / "big.pdf"
    big.write_bytes(_make_pdf("word " * 400))
    old_cap = FP._EXTRACT_INLINE
    FP._EXTRACT_INLINE = 50
    try:
        out_big = FP.file_processor({"file_path": str(big), "action": "extract_text"})
    finally:
        FP._EXTRACT_INLINE = old_cap
    check("pdf: truncated marker", "...[truncated]" in out_big, out_big[:200])
    check("pdf: truncated still has content", "word word" in out_big)
    check("pdf: truncated notes full size",
          "Full text (" in out_big and "chars) saved to" in out_big, out_big[-160:])

    # ── docx: content inline + file saved ───────────────────────────────────
    from docx import Document
    doc = Document()
    doc.add_paragraph("Name Md. Abdul Kader")
    doc.add_paragraph("Mobile 01717-288 187")
    docx_path = tmp / "resume.docx"
    doc.save(str(docx_path))
    out_d = FP.file_processor({"file_path": str(docx_path), "action": "extract_text"})
    check("docx: content inline", "Md. Abdul Kader" in out_d and
          "01717-288 187" in out_d, out_d[:200])
    check("docx: saved file", (tmp / "resume_extracted.txt").exists())

    # ── txt: unchanged behavior (returns content directly) ──────────────────
    txt = tmp / "note.txt"
    txt.write_text("hello from a text file", encoding="utf-8")
    out_t = FP.file_processor({"file_path": str(txt), "action": "extract_text"})
    check("txt: content inline", "hello from a text file" in out_t, out_t[:120])

    # ── pptx: content inline + file saved ───────────────────────────────────
    from pptx import Presentation
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "CV Phone 01717-288 187"
    pptx_path = tmp / "deck.pptx"
    prs.save(str(pptx_path))
    out_p = FP.file_processor({"file_path": str(pptx_path), "action": "extract_text"})
    check("pptx: content inline", "01717-288 187" in out_p, out_p[:200])
    check("pptx: saved file", (tmp / "deck_text.txt").exists())

    # ── TOOL dict still validates for discovery ─────────────────────────────
    rec = _validate(FP, "file_processor.py")
    check("tool: valid", rec.valid and rec.error == "", rec.error)
    check("tool: name", rec.name == "file_processor")

    print(f"\n{PASS} passed, {FAIL} failed")
    return 1 if FAIL else 0

if __name__ == "__main__":
    sys.exit(main())
