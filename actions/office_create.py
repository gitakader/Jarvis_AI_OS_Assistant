"""
office_create.py — Create Office documents from structured data.

Creates xlsx (Excel), docx (Word), and pptx (PowerPoint) files from structured
data provided by the model. All three types accept the same two input formats:

  A) headers + rows:  {"headers": [...], "rows": [[...], ...]}
  B) array of objects: [{"col": val, ...}, ...]

The path may be a full path, a shortcut like "desktop/report.xlsx", or omit
the extension (the correct suffix is appended automatically).
"""
from __future__ import annotations

import json
import os
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Union

_HAS_OPENPYXL = False
try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, numbers
    from openpyxl.utils import get_column_letter
    _HAS_OPENPYXL = True
except ImportError:
    pass

_HAS_DOCX = False
try:
    from docx import Document as _Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    _HAS_DOCX = True
except ImportError:
    pass

_HAS_PPTX = False
try:
    from pptx import Presentation as _Presentation
    from pptx.util import Inches as PptxInches, Pt as PptxPt, Emu
    from pptx.dml.color import RGBColor as PptxRgbColor
    from pptx.enum.text import PP_ALIGN
    _HAS_PPTX = True
except ImportError:
    pass

from actions.file_controller import _is_safe_path, _resolve_path

_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}")

_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_BODY_FONT   = Font(name="Calibri", size=11)
_THIN_BORDER = Border(
    left=Side(style="thin", color="D9E2F3"),
    right=Side(style="thin", color="D9E2F3"),
    top=Side(style="thin", color="D9E2F3"),
    bottom=Side(style="thin", color="D9E2F3"),
)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _maybe_date(val: Any) -> bool:
    if isinstance(val, (date, datetime)):
        return True
    if isinstance(val, str):
        return bool(_DATE_RE.match(val) or _DATETIME_RE.match(val))
    return False


def _auto_width(ws, min_width: int = 8, max_width: int = 40):
    """Set each column's width to a sensible value based on content."""
    for col_cells in ws.columns:
        best = min_width
        for cell in col_cells:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            if length > best:
                best = length
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(best + 2, max_width)


def _normalise_data(data: Any) -> tuple[list[str], list[list]]:
    """Convert the model's data into (headers, rows) no matter the input shape.

    Returns (headers, rows) or raises ValueError with a human message."""
    if data is None:
        raise ValueError("data is required (headers+rows or array of objects)")

    if isinstance(data, dict) and "headers" in data and "rows" in data:
        headers = [str(h) for h in data["headers"]]
        rows = data["rows"]
        if not isinstance(rows, list):
            raise ValueError("rows must be a list of lists")
        return headers, rows

    if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
        headers = list(data[0].keys())
        rows = [[row.get(h, "") for h in headers] for row in data]
        return headers, rows

    raise ValueError(
        "data must be either {'headers': [...], 'rows': [[...]]} "
        "or an array of objects like [{'col': val, ...}]"
    )


def _write_cell(ws, row: int, col: int, value: Any, is_header: bool = False):
    cell = ws.cell(row=row, column=col, value=value)
    if is_header:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
    else:
        cell.font = _BODY_FONT
        if isinstance(value, (int, float)):
            cell.alignment = Alignment(horizontal="right")
        elif isinstance(value, (date, datetime)):
            cell.number_format = "YYYY-MM-DD"
        elif isinstance(value, str) and _DATE_RE.match(value):
            cell.number_format = "YYYY-MM-DD"
    cell.border = _THIN_BORDER


def _build_xlsx(headers: list[str], rows: list[list], path: Path) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # headers
    for ci, h in enumerate(headers, start=1):
        _write_cell(ws, 1, ci, h, is_header=True)

    # data
    for ri, row in enumerate(rows, start=2):
        for ci, val in enumerate(row, start=1):
            _write_cell(ws, ri, ci, val, is_header=False)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    _auto_width(ws)

    os.makedirs(path.parent, exist_ok=True)
    wb.save(str(path))
    return (
        f"Spreadsheet created: {path} — "
        f"{len(rows)} row{'s' if len(rows) != 1 else ''} "
        f"× {len(headers)} column{'s' if len(headers) != 1 else ''}"
    )


_DOCX_HEADER_BG = "2F5496"
_DOCX_BORDER_COLOR = "D9E2F3"


def _set_cell_shading(cell, color: str):
    """Set background fill on a docx table cell via XML."""
    from lxml import etree
    shading = etree.SubElement(cell._tc.get_or_add_tcPr(), qn("w:shd"))
    shading.set(qn("w:fill"), color)
    shading.set(qn("w:val"), "clear")


def _set_cell_border(cell, **kwargs):
    """Set borders on a docx table cell. kwargs: top, bottom, left, right as dicts of {sz, val, color}."""
    from lxml import etree
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = etree.SubElement(tcPr, qn("w:tcBorders"))
    for edge in ("top", "bottom", "left", "right"):
        if edge in kwargs:
            attr = kwargs[edge]
            el = etree.SubElement(tcBorders, qn(f"w:{edge}"))
            el.set(qn("w:val"), attr["val"])
            el.set(qn("w:sz"), attr["sz"])
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), attr["color"])


def _build_docx(headers: list[str], rows: list[list], path: Path) -> str:
    doc = _Document()

    # Narrow margins for data tables
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"

    # header row
    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(h))
        run.font.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _set_cell_shading(cell, _DOCX_HEADER_BG)

    # data rows
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            p = cell.paragraphs[0]
            if isinstance(val, (int, float)):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            text = str(val) if val is not None else ""
            run = p.add_run(text)
            run.font.size = Pt(11)

    # thin borders on every cell
    border_attr = {"sz": "4", "val": "single", "color": _DOCX_BORDER_COLOR}
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(cell, top=border_attr, bottom=border_attr,
                             left=border_attr, right=border_attr)

    os.makedirs(path.parent, exist_ok=True)
    doc.save(str(path))
    return (
        f"Document created: {path} — "
        f"{len(rows)} row{'s' if len(rows) != 1 else ''} "
        f"× {len(headers)} column{'s' if len(headers) != 1 else ''}"
    )


_PPTX_HEADER_BG = PptxRgbColor(0x2F, 0x54, 0x96)
_PPTX_BORDER_COLOR = "D9E2F3"


def _pptx_table_borders(table):
    """Apply thin borders to every cell in a pptx table via XML."""
    from lxml import etree
    nsmap = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    stroke_color = _PPTX_BORDER_COLOR
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            # Remove existing borders
            for old in tcPr.findall("a:lnL", nsmap):
                tcPr.remove(old)
            for old in tcPr.findall("a:lnR", nsmap):
                tcPr.remove(old)
            for old in tcPr.findall("a:lnT", nsmap):
                tcPr.remove(old)
            for old in tcPr.findall("a:lnB", nsmap):
                tcPr.remove(old)
            for edge_tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, etree.QName(nsmap["a"], edge_tag[2:]))
                ln.set("w", "6350")  # 0.5pt in EMU
                solidFill = etree.SubElement(ln, etree.QName(nsmap["a"], "solidFill"))
                srgbClr = etree.SubElement(solidFill, etree.QName(nsmap["a"], "srgbClr"))
                srgbClr.set("val", stroke_color)


def _build_pptx(headers: list[str], rows: list[list], path: Path) -> str:
    prs = _Presentation()
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)

    # Blank slide
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    num_rows = 1 + len(rows)
    num_cols = len(headers)

    # Calculate column widths: auto-fit based on content
    col_max_chars = []
    for ci, h in enumerate(headers):
        best = len(str(h))
        for row in rows:
            val_len = len(str(row[ci])) if ci < len(row) else 0
            if val_len > best:
                best = val_len
        col_max_chars.append(best)

    total_chars = sum(col_max_chars) + num_cols  # padding
    available_emu = int(prs.slide_width * 0.92)  # 92% of slide width
    char_width_emu = max(1, available_emu // max(1, total_chars))

    table_left = PptxInches(0.5)
    table_top = PptxInches(0.5)
    table_width = available_emu
    table_height = min(PptxInches(6.0), (num_rows * PptxInches(0.4)))

    table_shape = slide.shapes.add_table(num_rows, num_cols, table_left, table_top, table_width, table_height)
    table = table_shape.table

    # Set column widths proportionally
    for ci in range(num_cols):
        table.columns[ci].width = col_max_chars[ci] * char_width_emu

    # Header row
    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(h)
        run.font.bold = True
        run.font.size = PptxPt(11)
        run.font.color.rgb = PptxRgbColor(0xFF, 0xFF, 0xFF)
        # Header fill
        tcPr = cell._tc.get_or_add_tcPr()
        from lxml import etree
        nsmap_a = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        solidFill = etree.SubElement(tcPr, etree.QName(nsmap_a["a"], "solidFill"))
        srgbClr = etree.SubElement(solidFill, etree.QName(nsmap_a["a"], "srgbClr"))
        srgbClr.set("val", "2F5496")

    # Data rows
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            if isinstance(val, (int, float)):
                p.alignment = PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = str(val) if val is not None else ""
            run.font.size = PptxPt(11)

    _pptx_table_borders(table)

    os.makedirs(path.parent, exist_ok=True)
    prs.save(str(path))
    return (
        f"Presentation created: {path} — "
        f"{len(rows)} row{'s' if len(rows) != 1 else ''} "
        f"× {len(headers)} column{'s' if len(headers) != 1 else ''}"
    )


# ── tool handler ─────────────────────────────────────────────────────────────


def office_create(parameters: dict, player=None, session_memory=None) -> str:
    params = parameters or {}
    doc_type = str(params.get("type") or "xlsx").lower().strip()
    raw_path = str(params.get("path") or "").strip()
    data     = params.get("data")

    if doc_type not in ("xlsx", "docx", "pptx"):
        return f"I can only create xlsx, docx, and pptx files right now — '{doc_type}' is not supported yet."

    if not raw_path:
        return "I need a path to save the file, sir — e.g. path 'desktop/report.xlsx'."

    if data is None:
        return "I need the data for the document — use the 'data' parameter with headers + rows or an array of objects."

    try:
        resolved = _resolve_path(raw_path)
    except Exception as e:
        return f"Could not resolve path '{raw_path}': {e}"

    expected_suffix = f".{doc_type}"
    if resolved.suffix.lower() != expected_suffix:
        resolved = resolved.with_suffix(expected_suffix)

    if not _is_safe_path(resolved):
        return f"Path '{resolved}' is outside accessible filesystem roots — I cannot write there."

    try:
        headers, rows = _normalise_data(data)
    except ValueError as e:
        return str(e)

    if not headers:
        return f"{doc_type.upper()} document must have at least one column (headers cannot be empty)."

    if doc_type == "xlsx":
        if not _HAS_OPENPYXL:
            return "openpyxl is not installed — cannot create .xlsx files."
        return _build_xlsx(headers, rows, resolved)

    if doc_type == "docx":
        if not _HAS_DOCX:
            return "python-docx is not installed — cannot create .docx files."
        return _build_docx(headers, rows, resolved)

    if not _HAS_PPTX:
        return "python-pptx is not installed — cannot create .pptx files."
    return _build_pptx(headers, rows, resolved)


TOOL = {
    "name": "office_create",
    "description": (
        "Creates an Office document from structured data. Supports .xlsx Excel "
        "spreadsheets, .docx Word documents, and .pptx PowerPoint presentations. "
        "Use this when the user asks to 'create a spreadsheet', 'make a budget in "
        "Excel', 'write a Word document', 'generate a report as docx', 'create a "
        "PowerPoint', 'make a presentation', and similar. "
        "Pass the document data as structured JSON — either headers+rows or an "
        "array of objects. NOT for reading existing documents (use file_processor "
        "extract_text or csv_reader for that)."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "type": {
                "type": "STRING",
                "description": "Document type: 'xlsx' for Excel, 'docx' for Word, 'pptx' for PowerPoint.",
            },
            "path": {
                "type": "STRING",
                "description": (
                    "Where to save the file — full path or shortcut like "
                    "'desktop/report.xlsx' (extension appended automatically if missing)."
                ),
            },
            "data": {
                "type": "OBJECT",
                "description": (
                    "The document content. Two accepted shapes:\n"
                    "A) headers + rows: {'headers': ['Name','Age'], 'rows': [['Alice',30],['Bob',25]]}\n"
                    "B) array of objects: [{'Name':'Alice','Age':30},{'Name':'Bob','Age':25}]"
                ),
            },
        },
        "required": ["type", "path", "data"],
    },
    "handler": office_create,
}